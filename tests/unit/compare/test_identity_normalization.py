"""Object identity across databases (ADR 0007): case-insensitive names,
the snapshot's default schema counts as empty, colliding names inside one
database match exactly, and findings carry every original name."""

from dataclasses import replace

from compare.conftest import make_column, make_table

from schema_comparator.compare.engine import compare_snapshots
from schema_comparator.compare.models import (
    ColumnMismatch,
    MissingColumn,
    MissingProcedure,
    MissingTable,
)
from schema_comparator.domain.comparison.models import column_label, object_label
from schema_comparator.domain.schema.models import ProcedureSnapshot, SchemaSnapshot


def _snapshot(name, *tables, default_schema=None, provider="sqlserver", procs=()):
    return SchemaSnapshot(
        profile_name=name,
        provider_id=provider,
        tables=tuple(tables),
        procedures=tuple(procs),
        default_schema=default_schema,
    )


def test_mysql_databases_with_different_names_match_their_tables() -> None:
    dev = _snapshot(
        "dev",
        make_table("app_dev", "poliza", make_column("id")),
        default_schema="app_dev",
        provider="mysql",
    )
    prod = _snapshot(
        "prod",
        make_table("app_prod", "poliza", make_column("id")),
        default_schema="app_prod",
        provider="mysql",
    )

    assert compare_snapshots([dev, prod]).entries == ()


def test_oracle_default_owner_is_normalized_but_other_owners_are_not() -> None:
    dev = _snapshot(
        "dev",
        make_table("APP_DEV", "POLIZA", make_column("ID")),
        make_table("SHARED", "TARIFA", make_column("ID")),
        default_schema="APP_DEV",
        provider="oracle",
    )
    prod = _snapshot(
        "prod",
        make_table("APP_PROD", "POLIZA", make_column("ID")),
        make_table("OTHER", "TARIFA", make_column("ID")),
        default_schema="APP_PROD",
        provider="oracle",
    )

    missing = [e for e in compare_snapshots([dev, prod]).entries]

    assert {(e.qualified_name, e.missing_from_profile) for e in missing} == {
        (("OTHER", "TARIFA"), "dev"),
        (("SHARED", "TARIFA"), "prod"),
    }


def test_table_and_column_names_match_ignoring_case_and_report_all_variants() -> None:
    a = _snapshot(
        "a", make_table("dbo", "Users", make_column("Email", data_type="varchar"))
    )
    b = _snapshot(
        "b", make_table("dbo", "users", make_column("email", data_type="nvarchar"))
    )

    entries = compare_snapshots([a, b]).entries

    assert len(entries) == 1
    mismatch = entries[0]
    assert isinstance(mismatch, ColumnMismatch)
    assert mismatch.qualified_name == ("dbo", "Users")
    assert mismatch.column_name == "Email"
    assert object_label(mismatch) == "dbo.Users / dbo.users"
    assert column_label(mismatch) == "Email / email"


def test_identical_names_have_no_variants() -> None:
    a = _snapshot("a", make_table("dbo", "Users", make_column("id")))
    b = _snapshot("b", make_table("dbo", "Users"))

    (entry,) = compare_snapshots([a, b]).entries

    assert isinstance(entry, MissingColumn)
    assert object_label(entry) == "dbo.Users"
    assert column_label(entry) == "id"


def test_names_colliding_inside_one_database_match_exactly() -> None:
    a = _snapshot(
        "a",
        make_table("public", "Users", make_column("id")),
        make_table("public", "users", make_column("id")),
        provider="postgresql",
    )
    b = _snapshot(
        "b", make_table("public", "users", make_column("id")), provider="postgresql"
    )

    entries = compare_snapshots([a, b]).entries

    assert [(type(e), e.qualified_name, e.missing_from_profile) for e in entries] == [
        (MissingTable, ("public", "Users"), "b")
    ]


def test_routines_match_ignoring_case() -> None:
    a = _snapshot("a", procs=(ProcedureSnapshot("dbo", "usp_Calcular"),))
    b = _snapshot("b", procs=(ProcedureSnapshot("dbo", "USP_CALCULAR"),))
    c = _snapshot("c")

    entries = compare_snapshots([a, b, c]).entries

    assert [type(e) for e in entries] == [MissingProcedure]
    assert entries[0].missing_from_profile == "c"
    assert object_label(entries[0]) == "dbo.usp_Calcular / dbo.USP_CALCULAR"


def test_missing_table_ddl_uses_canonical_column_names() -> None:
    a = _snapshot("a", make_table("dbo", "Users", make_column("Email")))
    b = _snapshot("b", make_table("dbo", "users", make_column("email")))
    c = _snapshot("c")

    (entry,) = compare_snapshots([a, b, c]).entries

    assert isinstance(entry, MissingTable)
    assert entry.table_name == "Users"
    assert [col.name for _, cols in entry.present_columns for col in cols] == [
        "Email",
        "Email",
    ]


def test_snapshot_default_schema_defaults_to_none() -> None:
    snapshot = SchemaSnapshot(profile_name="x", tables=())

    assert snapshot.default_schema is None
    assert replace(snapshot, default_schema="app").default_schema == "app"


def _case_variant_result():
    a = _snapshot(
        "a", make_table("dbo", "Users", make_column("Email", data_type="varchar"))
    )
    b = _snapshot(
        "b", make_table("dbo", "users", make_column("email", data_type="nvarchar"))
    )
    return compare_snapshots([a, b])


def test_console_html_excel_and_tui_show_every_spelling() -> None:
    import io

    from openpyxl import load_workbook

    from schema_comparator.report.console import render_console
    from schema_comparator.report.excel import export_excel
    from schema_comparator.report.html import render_html
    from schema_comparator.tui.formatting import build_tree_data, detail_text

    result = _case_variant_result()

    console = render_console(result)
    assert "dbo.Users / dbo.users: 1 hallazgo(s)" in console
    assert "[Email / email]" in console

    html = render_html(result)
    assert "dbo.Users / dbo.users" in html
    assert "Email / email" in html

    sheet = load_workbook(io.BytesIO(export_excel(result))).active
    values = [cell for row in sheet.iter_rows(values_only=True) for cell in row]
    assert "dbo.Users / dbo.users" in values
    assert "Email / email" in values

    (group,) = build_tree_data(result).groups
    assert group.qualified_label == "dbo.Users / dbo.users"
    assert detail_text(result.entries[0]).startswith(
        "Columna: dbo.Users / dbo.users → Email / email"
    )
